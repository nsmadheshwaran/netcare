"""Render a Report (title, columns, rows, totals, notes) as CSV, Excel or PDF."""
import csv
import io
from xml.sax.saxutils import escape as esc
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle


@dataclass
class Column:
    key: str
    label: str
    kind: str = "text"  # text | money | qty | int | date | pct


@dataclass
class Report:
    title: str
    columns: list[Column]
    rows: list[dict]
    subtitle: str = ""
    totals: dict | None = None
    notes: list[str] = field(default_factory=list)
    draft: bool = False  # tax summaries etc: clearly not for filing
    sections: list["Report"] = field(default_factory=list)

    def to_json(self) -> dict:
        def conv(v):
            return str(v) if isinstance(v, (Decimal, date, datetime)) else v
        return {"title": self.title, "subtitle": self.subtitle, "draft": self.draft, "notes": self.notes,
                "columns": [c.__dict__ for c in self.columns],
                "rows": [{k: conv(v) for k, v in r.items()} for r in self.rows],
                "totals": {k: conv(v) for k, v in (self.totals or {}).items()} or None,
                "sections": [s.to_json() for s in self.sections]}


def _safe(v):
    """Neutralise spreadsheet formula injection in text cells."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") and not _is_number(v):
        return "'" + v
    return v


def _is_number(s: str) -> bool:
    try:
        Decimal(s)
        return True
    except Exception:
        return False


def _all(report: Report) -> list[Report]:
    return [report, *report.sections]


def to_csv(report: Report) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([report.title + (" (DRAFT - NOT FOR FILING)" if report.draft else "")])
    if report.subtitle:
        w.writerow([report.subtitle])
    for part in _all(report):
        if part is not report:
            w.writerow([])
            w.writerow([part.title])
        w.writerow([c.label for c in part.columns])
        for r in part.rows:
            w.writerow([_safe(r.get(c.key, "")) for c in part.columns])
        if part.totals:
            w.writerow([_safe(part.totals.get(c.key, "")) for c in part.columns])
    for n in report.notes:
        w.writerow([])
        w.writerow([f"Note: {n}"])
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens the rupee sign correctly


def to_xlsx(report: Report) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = report.title[:31]
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E0E7FF")
    ws.append([report.title + (" (DRAFT - NOT FOR FILING)" if report.draft else "")])
    ws["A1"].font = Font(bold=True, size=13)
    if report.subtitle:
        ws.append([report.subtitle])
    widths: dict[int, int] = {}
    for part in _all(report):
        ws.append([])
        if part is not report:
            ws.append([part.title])
            ws.cell(ws.max_row, 1).font = bold
        ws.append([c.label for c in part.columns])
        for i in range(1, len(part.columns) + 1):
            ws.cell(ws.max_row, i).font = bold
            ws.cell(ws.max_row, i).fill = head_fill
        rows = list(part.rows) + ([part.totals] if part.totals else [])
        for ri, r in enumerate(rows):
            vals = []
            for c in part.columns:
                v = r.get(c.key, "")
                if c.kind in ("money", "qty", "pct") and v not in ("", None):
                    v = float(Decimal(str(v)))  # Excel stores numbers as floats; values are exact to 2-3 dp
                vals.append(_safe(v))
            ws.append(vals)
            for i, c in enumerate(part.columns, start=1):
                cell = ws.cell(ws.max_row, i)
                if c.kind == "money":
                    cell.number_format = '#,##0.00'
                elif c.kind == "qty":
                    cell.number_format = '#,##0.###'
                if part.totals and ri == len(rows) - 1:
                    cell.font = bold
                widths[i] = max(widths.get(i, 8), min(45, len(str(cell.value or "")) + 2))
    for n in report.notes:
        ws.append([])
        ws.append([f"Note: {n}"])
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=False)
    for i, wdt in widths.items():
        ws.column_dimensions[ws.cell(1, i).column_letter].width = wdt
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def fmt(v, kind: str) -> str:
    if v in ("", None):
        return ""
    if kind == "money":
        return inr(v)
    if kind == "qty":
        return f"{Decimal(str(v)).normalize():f}"
    return str(v)


def inr(v) -> str:
    """Indian digit grouping: 12,34,567.89"""
    d = Decimal(str(v)).quantize(Decimal("0.01"))
    neg = d < 0
    whole, frac = f"{abs(d):.2f}".split(".")
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        groups = []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        if head:
            groups.insert(0, head)
        whole = ",".join(groups + [tail])
    return f"{'-' if neg else ''}{whole}.{frac}"


def to_pdf(report: Report, org_name: str) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=report.title, author=org_name)
    ss = getSampleStyleSheet()
    small = ss["BodyText"].clone("small", fontSize=7.5, leading=9)
    num = small.clone("num", alignment=2)
    numeric = ("money", "qty", "int", "pct")

    def cell(text, c, bold=False):
        t = esc(text)
        return Paragraph(f"<b>{t}</b>" if bold else t, num if c.kind in numeric else small)
    story = [Paragraph(esc(org_name), ss["Heading3"]), Paragraph(esc(report.title), ss["Title"])]
    if report.draft:
        story.append(Paragraph("<font color='#b91c1c'><b>DRAFT, NOT FOR FILING.</b> Prepared from your records "
                               "for review by a qualified accountant.</font>", ss["BodyText"]))
    if report.subtitle:
        story.append(Paragraph(esc(report.subtitle), ss["BodyText"]))
    for part in _all(report):
        story.append(Spacer(1, 4 * mm))
        if part is not report:
            story.append(Paragraph(esc(part.title), ss["Heading4"]))
        data = [[Paragraph(f"<b>{esc(c.label)}</b>", small) for c in part.columns]]
        for r in part.rows:
            data.append([cell(fmt(r.get(c.key, ""), c.kind), c, r.get("bold", False)) for c in part.columns])
        if part.totals:
            data.append([cell(fmt(part.totals.get(c.key, ""), c.kind), c, True) for c in part.columns])
        if len(data) == 1:
            data.append([Paragraph("No records in this period.", small)] + [""] * (len(part.columns) - 1))
        t = LongTable(data, repeatRows=1)
        style = [("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                 ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e0e7ff")),
                 ("VALIGN", (0, 0), (-1, -1), "TOP")]
        t.setStyle(TableStyle(style))
        story.append(t)
    for n in report.notes:
        story.append(Spacer(1, 2 * mm))
        story.append(Paragraph(f"Note: {esc(n)}", small))
    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph(f"Generated {datetime.now().strftime('%d-%m-%Y %H:%M')} by NetCare Business Suite",
                           small))
    doc.build(story)
    return buf.getvalue()
